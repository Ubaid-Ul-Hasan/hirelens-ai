import pytest

from src.ingestion.document_ingestion import ingest_document
from src.utils.validation import ValidationError


def test_txt_ingestion_happy_path():
    result = ingest_document(b"Experienced Python developer with 5 years in ML.", "resume.txt")
    assert result.quality.is_empty is False
    assert "Python" in result.raw_text


def test_rejects_unsupported_extension():
    with pytest.raises(ValidationError):
        ingest_document(b"data", "resume.exe")


def test_rejects_empty_file():
    with pytest.raises(ValidationError):
        ingest_document(b"", "resume.txt")


def test_rejects_oversized_file():
    from src.utils.config import CONFIG

    max_bytes = CONFIG.ingestion.max_file_size_mb * 1024 * 1024
    with pytest.raises(ValidationError):
        ingest_document(b"x" * (max_bytes + 1), "resume.txt", size_bytes=max_bytes + 1)


def test_short_content_flagged_but_not_rejected():
    result = ingest_document(b"hi", "resume.txt")
    assert result.quality.is_empty is False
    assert result.quality.message is not None  # warned, not raised


def test_docx_ingestion_extracts_paragraphs_and_tables():
    import io

    from docx import Document

    doc = Document()
    doc.add_paragraph("Jane Doe")
    doc.add_paragraph("Senior ML Engineer")
    table = doc.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Python"
    table.rows[0].cells[1].text = "PyTorch"
    buf = io.BytesIO()
    doc.save(buf)

    result = ingest_document(buf.getvalue(), "resume.docx")
    assert "Jane Doe" in result.raw_text
    assert "Python" in result.raw_text
    assert "PyTorch" in result.raw_text
